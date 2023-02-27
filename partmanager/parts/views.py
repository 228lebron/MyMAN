import openpyxl
import matplotlib.pyplot as plt
import matplotlib
matplotlib.use('agg')
from openpyxl.writer.excel import save_virtual_workbook
import pandas as pd
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.core.cache import cache
from django.views import View
from django.utils import timezone
from django.contrib.auth.decorators import user_passes_test, login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.contrib.auth import authenticate, login
from django.contrib.auth.forms import AuthenticationForm
from django.core.exceptions import ValidationError
from django.db.models import Q, F, QuerySet, Min, Subquery
from django.contrib.auth.models import User
from django.urls import reverse_lazy
from django.views.generic import CreateView, TemplateView, UpdateView
from .models import Request, Quota, Part, RequestQuotaResult
from .filters import ReqFilter, QuotaFilter, PartFilter, RequestFilter
from .forms import QuotaUploadForm, PartForm, RequestQuotaForm
from .tables import PartTable, RequestTable, RequestQuotaTable
from django_tables2 import RequestConfig
from django.db.models import Count
from django.core.paginator import Paginator
from functools import wraps


def create_quotas_from_xlsx(file):
    wb = openpyxl.load_workbook(file)
    ws = wb.active

    quotas = []
    for row in ws.iter_rows(values_only=True):
        part_number, brand, quantity, price, datecode, lead_time, supplier, date = row
        try:
            part = Part.objects.get(number=part_number, brand=brand)
        except Part.DoesNotExist:
            part = Part.objects.create(number=part_number, series=f"Введите серию {part_number[:6]}", brand=brand)
        quotas.append(Quota(part=part, quantity=quantity, price=price, datecode=datecode, supplier=supplier,
                            lead_time=lead_time, date=date))
    Quota.objects.bulk_create(quotas)

def group_required(group_name):
    """
    Decorator that checks if the user is in the specified group.
    """
    def check_group(user):
        return user.groups.filter(name=group_name).exists()

    return user_passes_test(check_group)

# Initial
def initial_view(request):
    if request.user.groups.filter(name='sale').exists():
        # Redirect sales managers to the sales manager view
        return redirect('responses/')
    elif request.user.groups.filter(name='product').exists():
        # Redirect product managers to the product manager view
        return redirect('products/list/')
    else:
        return redirect('accounts/login/')

####################### Base Views #####################################
class BaseView(View):
    def dispatch(self, request, *args, **kwargs):
        current_user_groups = request.user.groups.values_list("name", flat=True)
        context = {
            "is_sale": "sale" in current_user_groups,
            "is_product_man": "product" in current_user_groups,
        }
        response = super().dispatch(request, *args, **kwargs)
        if hasattr(response, 'context_data'):
            response.context_data.update(context)
        else:
            response.context = context
        return response
########################################################################


# Product manager
########################################################################
#@login_required()
#@group_required('product')
#def product_manager_view(request):
#    current_user_groups = request.user.groups.values_list("name", flat=True)
#    result_qs = RequestQuotaResult.objects.all().order_by('-request_id')
#    myFilter = ReqFilter(request.GET, queryset=result_qs)
#    result_qs = myFilter.qs
#    table = RequestQuotaTable(result_qs)
#    RequestConfig(request).configure(table)
#    context = {
#        'result': result_qs,
#        'filter': myFilter,
#        "is_sale": "sale" in current_user_groups,
#        "is_product_man": "product" in current_user_groups,
#        'table': table,
#    }
#
#    return render(request, 'product_view.html', context)


class ProductManagerView(UserPassesTestMixin, BaseView):
    def test_func(self):
        return self.request.user.groups.filter(name='product').exists()

    def get(self, request, *args, **kwargs):
        result_qs = RequestQuotaResult.objects.all().order_by('-request_id')
        myFilter = ReqFilter(request.GET, queryset=result_qs)
        result_qs = myFilter.qs
        table = RequestQuotaTable(result_qs)
        RequestConfig(request).configure(table)
        context = {
            'result': result_qs,
            'filter': myFilter,
            'table': table,
        }
        return render(request, 'product_view.html', context)

@login_required()
@group_required('product')
def upload_quotas(request):
    if request.method == 'POST':
        form = QuotaUploadForm(request.POST, request.FILES)
        if form.is_valid():
            create_quotas_from_xlsx(request.FILES['file'])
            return redirect('/quotas/')
    else:
        form = QuotaUploadForm()
    return render(request, 'quotas_upload.html', {'form': form})




# Sale Manager views
###################################################
@login_required()
@group_required('sale')
def sale_manager_view(request):
    current_user_groups = request.user.groups.values_list("name", flat=True)
    manager = request.user
    requests = Request.objects.filter(manager=manager).order_by('-date')
    request_filter = RequestFilter(request.GET, queryset=requests)
    requests = request_filter.qs
    context = {
        'requests': requests,
        'filter': request_filter,
        'user': manager,
        "is_sale": "sale" in current_user_groups,
        "is_product": "product" in current_user_groups,
    }
    return render(request, 'sales_manager/manager_view.html', context)

@login_required()
def quota_list(request):
    current_user_groups = request.user.groups.values_list("name", flat=True)
    if "sale" in current_user_groups:
        return redirect('responses/')

    result_qs = Quota.objects.all().order_by('id')

    quotaFilter = QuotaFilter(request.GET, queryset=result_qs)
    result_qs = quotaFilter.qs

    return render(request, 'quotas_list.html', {'result': result_qs, 'filter': quotaFilter})

@login_required()
def part_list(request):
    current_user_groups = request.user.groups.values_list("name", flat=True)
    parts = Part.objects.all()
    partFilter = PartFilter(request.GET, queryset=parts)
    parts = partFilter.qs
    table = PartTable(parts)
    RequestConfig(request).configure(table)
    context = {
        'parts': parts,
        'filter': partFilter,
        'table': table,
        "is_sale": "sale" in current_user_groups,
        "is_product_man": "product" in current_user_groups,
    }
    return render(request, 'parts/part_list.html', context)

def get_price_data(part):
    quotas = Quota.objects.filter(part=part)
    prices = []
    dates = []
    for quota in quotas:
        prices.append(quota.price)
        dates.append(quota.date)
    return prices, dates
class PartDetailView(LoginRequiredMixin, TemplateView):
    template_name = 'parts/part_detail.html'
    login_url = 'accounts/login/'

    def get(self, request, *args, **kwargs):
        part_id = kwargs.get('part_id')
        part = get_object_or_404(Part, pk=part_id)
        quotas = Quota.objects.filter(part=part).order_by('-date')
        requests = Request.objects.filter(part=part)
        total_requests = requests.count()
        last_quota_date = quotas.last().date if quotas.exists() else None
        last_price = quotas.last().price if quotas.exists() else None

        prices, dates = get_price_data(part)

        # generate the price chart
        fig, ax = plt.subplots()
        ax.tick_params(axis='x', labelsize=6)
        ax.plot(dates, prices)
        ax.set(xlabel='Дата', ylabel='Цена USD', title='График цен для {}'.format(part.number))
        ax.grid()

        # convert the plot to a PNG image and add it to the context
        from io import BytesIO
        import base64
        buffer = BytesIO()
        plt.savefig(buffer, format='png')
        buffer.seek(0)
        image_png = buffer.getvalue()
        buffer.close()
        graph = base64.b64encode(image_png).decode('utf-8')

        return render(request, self.template_name, {'quotas': quotas, 'part': part, 'total_requests': total_requests,
                                                    'last_price': last_price, 'last_quota_date': last_quota_date,
                                                    'graph': graph})


class PartCreateView(CreateView, LoginRequiredMixin):
    model = Part
    fields = ['series', 'number', 'brand', 'case_type']
    template_name = 'parts/part_form.html'
    success_url = reverse_lazy('parts:part_list')
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        current_user_groups = self.request.user.groups.values_list("name", flat=True)
        context['is_sale'] = "sale" in current_user_groups
        context['is_product_man'] = "product" in current_user_groups
        return context


class PartUpdateView(UpdateView, LoginRequiredMixin):
    model = Part
    fields = ['series', 'number', 'brand', 'case_type']
    template_name = 'parts/part_edit.html'
    success_url = reverse_lazy('parts:part_list')
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        current_user_groups = self.request.user.groups.values_list("name", flat=True)
        context['is_sale'] = "sale" in current_user_groups
        context['is_product_man'] = "product" in current_user_groups
        return context


class RequestCreateView(CreateView, LoginRequiredMixin):
    model = Request
    fields = ['part', 'quantity', 'customer', 'date']
    template_name = 'requests/request_form.html'
    success_url = reverse_lazy('parts:initial_view')
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        current_user_groups = self.request.user.groups.values_list("name", flat=True)
        context['is_sale'] = "sale" in current_user_groups
        context['is_product_man'] = "product" in current_user_groups
        context['customers'] = Request.objects.values('customer').distinct()
        return context

    def form_valid(self, form):
        form.instance.manager = self.request.user
        return super().form_valid(form)

@login_required()
def request_list(request):
    current_user_groups = request.user.groups.values_list("name", flat=True)
    if "sale" in current_user_groups:
        return redirect('responses/')
    manager = request.user
    requests = Request.objects.all()
    reqFilter = RequestFilter(request.GET, queryset=requests)
    requests = reqFilter.qs
    table = RequestTable(requests)
    RequestConfig(request).configure(table)
    table.paginate(page=request.GET.get('page', 1), per_page=10)
    return render(request, 'requests/request_list.html', {'requests': requests, 'filter':reqFilter, 'user': manager, 'table': table, })
class AttachQuotaToRequestView(UpdateView):
    model = Request
    fields = ['selected_quota', 'currency', 'customer_price']
    template_name = 'attach_quota.html'
    success_url = reverse_lazy('parts:request_list')
    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        request = self.get_object()
        creation_date_cutoff = timezone.now() - timezone.timedelta(days=5)
        form.fields['selected_quota'].queryset = Quota.objects.filter(part__series=request.part.series,
                                                                      date__gte=creation_date_cutoff,)
        return form

    def form_valid(self, form):
        request = form.save(commit=False)
        request.status = 'QUOTA'  # change the status to 'quota'
        request.save()
        return super().form_valid(form)

def clients(request):
    current_user_groups = request.user.groups.values_list("name", flat=True)
    clients = Request.objects.values('customer').annotate(
        total_requests=Count('id', distinct=True),
        paid_requests=Count('id', filter=Q(status='PAID'), distinct=True),
        redemption_percent=(100 * Count('id', filter=Q(status='PAID')) / Count('id')),
    )
    context = {
        'clients': clients,
        "is_sale": "sale" in current_user_groups,
        "is_product_man": "product" in current_user_groups,
    }

    return render(request, 'clients/clients_list.html', context)


class AddAccountView(View):
    def post(self, request, pk):
        request_obj = get_object_or_404(Request, pk=pk)
        request_obj.status = 'PAID'
        request_obj.save()
        return redirect('/responses/')
